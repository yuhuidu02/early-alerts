"""
early_alerts/research_attribution.py

ON-DEMAND RESEARCH SCRIPT — separate from the live Chiron pipeline.

Purpose:
  For every student in the given course(s), run the attribution agent
  TWICE:
    - AI-only:  agent sees fired_rules (as context, not a constraint) +
                all 4 tools. No ML information.
    - AI+ML:    identical prompt, PLUS the ML model's at-risk probability.

  Unlike the live pipeline, the agent is NOT constrained to pick from
  fired_rules — it can choose freely among all 6 categories. This tests
  whether the agent's own judgment (with vs. without the ML signal)
  diverges from the deterministic rule engine, and whether the
  rationale text itself surfaces anything useful.

Does NOT:
  - write to the `alerts` table
  - write to / update `chiron_runs`
  - touch anything the instructor dashboard (alerts.js) reads from

Output:
  A single xlsx per invocation, written to:
    s3://<bucket>/research/YYYY/MM/DD/research_alerts_<timestamp>.xlsx

Usage (script or notebook, not wired into batch_handler):
    from research_attribution import run_research
    run_research(mydb_cfg, timescale_cfg, canvas_course_ids=[201124],
                  s3_bucket="early-alerts", as_of_date="2026-09-08")
"""

import json
import logging
import os
import re
from collections import defaultdict
from datetime import datetime

import boto3
import openpyxl
import psycopg2
import psycopg2.extras
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from strands import Agent
from strands.models import BedrockModel

from batch_categorize import (
    CAT_POSITIVE_SATISFACTORY,
    CAT_POSITIVE_EXCEPTIONAL,
    CAT_MISSING_ASSIGNMENTS,
    CAT_EXAM_PERFORMANCE,
    CAT_LOW_ENGAGEMENT,
    CAT_NEVER_ATTENDED,
    fetch_student_roster,
    fetch_click_signals,
    fetch_model_predictions,
    get_fired_rules,
)
from tools.attribution_tools import AttributionToolkit

logger = logging.getLogger()
logger.setLevel(logging.INFO)

# The 6 categories the agent may freely choose among (no CAT_NEEDS_AGENT here —
# that's a routing placeholder, never a real answer).
VALID_CATEGORIES = [
    CAT_NEVER_ATTENDED,
    CAT_MISSING_ASSIGNMENTS,
    CAT_LOW_ENGAGEMENT,
    CAT_EXAM_PERFORMANCE,
    CAT_POSITIVE_EXCEPTIONAL,
    CAT_POSITIVE_SATISFACTORY,
]

FALLBACK_CATEGORY = CAT_POSITIVE_SATISFACTORY  # used only if parsing totally fails and no fired_rules exist

# --- System prompt: free choice among all 6 categories, with criteria spelled out ---
SYSTEM_PROMPT = """
You are a student alert categorization agent at a university, running in RESEARCH MODE.

Unlike normal operation, you are NOT restricted to a pre-filtered list of categories.
Investigate using all 4 tools, then choose the ONE category that best describes this
student's current status from the full set below, using your own independent judgment.

CATEGORIES (choose exactly one):
- "Neg: Never Attended / No WebCampus Activity" — no meaningful Canvas/WebCampus activity recorded.
- "Neg: Missing Assignment(s)" — one or more missing assignments is the primary issue.
- "Neg: Lack of Engagement or Infrequent Attendance" — activity is present but sparse, shallow,
  or infrequent relative to peers (low click breadth/intensity/coherence).
- "Neg: Exam/Quiz Performance" — the primary issue is poor quiz/exam scores, not attendance or engagement.
- "Positive: Satisfactory Course Performance" — no significant negative signal; performing adequately.
- "Positive: Exceptional Course Performance" — no significant negative signal; performing strongly
  (current course grade around 85% or higher is a good rule of thumb, but use overall judgment).

You will also be told which categories the deterministic rule engine flagged for this student
("fired_rules"). Treat this as CONTEXT ONLY, not a constraint — you may agree with it, pick a
different category among the 6, or land on a positive category even if a negative rule fired,
if your investigation supports that.

RULES:
- Call ALL 4 tools before deciding.
- Pick the category that best explains the root cause, not just a symptom.
- Base your decision on the tool evidence, not just on which rules fired.

CONSTRUCT SCORING RULES:
- con (Confidence):  1-6, HIGHER is better.
- sth (Study Habits): 1-6, HIGHER is better.
- mot (Motivation):   1-6, HIGHER is better.
- abur (Stress):      1-6, HIGHER means MORE stressed.
- res (Resilience):   1-6, LOWER means LESS resilient.

GRADE FIELDS (from get_missing_assignments):
- latest_current_grade: the student's overall course grade (0-100). Context only.
- latest_quiz_score: the student's most recent Chiron quiz score (0-100).
- latest_missing_assignments: the student's most recent missing assignment count.
- history: full weekly snapshot ordered oldest to newest. Use the trend to distinguish
  chronic issues from sudden drops or improving trajectories.

Click z-scores: 0 = course average. Below -1.0 = warning. Below -2.0 = strong warning.

You may also be given an ML model's estimated probability of subsequent academic
difficulty. If given, weigh it as ONE signal among the others — not an automatic
override of your own judgment from the tool evidence.

OUTPUT FORMAT — return ONLY a JSON object, no markdown fences, no extra text:
{
  "category": "<exact category string from the list above>",
  "rationale": "<2-4 sentences citing the specific evidence that drove this decision>"
}
"""

# --- Config helpers (mirrors attribution_handler.py) ---
def _get_secrets(name: str) -> str:
    return boto3.client("secretsmanager", region_name=os.environ["AWS_DEFAULT_REGION"]) \
                .get_secret_value(SecretId=name)["SecretString"]

def _build_db_configs() -> tuple:
    password = _get_secrets(os.environ["DB_SECRET_NAME"])
    base = {"user": "dbuser", "password": password, "sslmode": "prefer", "connect_timeout": 10}
    return (
        {**base, "host": os.environ["MYDB_HOST"],
                 "port": int(os.environ.get("MYDB_PORT", 5432)),
                 "dbname": os.environ["MYDB_NAME"]},
        {**base, "host": os.environ["TIMESCALE_HOST"],
                 "port": int(os.environ.get("TIMESCALE_PORT", 5432)),
                 "dbname": os.environ["TIMESCALE_NAME"]},
    )

def _query(cfg: dict, sql: str, params=None) -> list:
    with psycopg2.connect(**cfg) as conn:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(sql, params)
            return [dict(row) for row in cur.fetchall()]

# --- Resolve timescale (click DB) internal ids in bulk ---
def resolve_timescale_ids(timescale_cfg: dict, canvas_user_ids: list, canvas_course_ids: list) -> tuple:
    """
    Returns (student_map, course_map):
      student_map: canvas_user_id -> timescale internal student id
      course_map:  canvas_course_id -> timescale internal course id
    """
    ph_u = ",".join(["%s"] * len(canvas_user_ids))
    ph_c = ",".join(["%s"] * len(canvas_course_ids))

    student_rows = _query(timescale_cfg,
        f"SELECT canvas_user_id, id FROM students WHERE canvas_user_id IN ({ph_u})",
        canvas_user_ids)
    course_rows = _query(timescale_cfg,
        f"SELECT canvas_course_id, id FROM courses WHERE canvas_course_id IN ({ph_c})",
        canvas_course_ids)

    student_map = {r["canvas_user_id"]: r["id"] for r in student_rows}
    course_map = {r["canvas_course_id"]: r["id"] for r in course_rows}
    return student_map, course_map

# --- Prompt building ---
def _build_prompt(student: dict, fired_rules: list, as_of_date: str, ids: dict,
                   ml_prediction: dict | None) -> str:
    fired_text = ", ".join(fired_rules) if fired_rules else "(none — no deterministic rule fired)"

    ml_block = ""
    if ml_prediction is not None:
        prob = ml_prediction.get("at_risk_probability")
        predicted = ml_prediction.get("at_risk_predicted")
        ml_block = f"""
          ML Model Risk Estimate:
          - Probability of subsequent academic difficulty: {prob if prob is not None else "unknown"}
          - Model's binary at-risk prediction: {predicted}
        """

    return f"""
      Student Name: {student['name']}
      Course: {student['course_name']}
      As of Date: {as_of_date}

      Deterministic rule engine flagged (context only, not a constraint): {fired_text}
      {ml_block}
      Internal IDs for tool calls:
      - Student ID (mydb):      {ids['mydb_student_id']}
      - Course ID (mydb):       {ids['mydb_course_id']}
      - Student ID (timescale): {ids['ts_student_id']}
      - Course ID (timescale):  {ids['ts_course_id']}

      Use all 4 tools, then return your decision as the JSON object described in your instructions.
    """

# --- Parsing the agent's structured output ---
_JSON_BLOCK_RE = re.compile(r"\{.*\}", re.DOTALL)

def _parse_agent_output(response_text: str, fired_rules: list) -> tuple:
    """
    Returns (category, rationale). Falls back gracefully if the agent
    didn't return valid JSON or returned an unrecognized category.
    """
    match = _JSON_BLOCK_RE.search(response_text)
    if match:
        try:
            parsed = json.loads(match.group(0))
            category = parsed.get("category")
            rationale = parsed.get("rationale", "").strip()
            if category in VALID_CATEGORIES and rationale:
                return category, rationale
        except (json.JSONDecodeError, AttributeError):
            pass

    # Fallback 1: category name appears somewhere in the raw text
    for cat in VALID_CATEGORIES:
        if cat in response_text:
            logger.warning("Could not parse structured JSON; recovered category by substring match: %s", cat)
            return cat, f"[Could not parse structured rationale — raw agent output: {response_text[:500]}]"

    # Fallback 2: total failure — default to first fired rule, or a positive default
    default_category = fired_rules[0] if fired_rules else FALLBACK_CATEGORY
    logger.warning("Agent output completely unparseable; defaulting to %s", default_category)
    return default_category, f"[Agent output unparseable — defaulted. Raw output: {response_text[:500]}]"

# --- Run a single agent condition (AI-only or AI+ML) ---
def _run_agent_condition(toolkit: AttributionToolkit, model: BedrockModel, student: dict,
                          fired_rules: list, as_of_date: str, ids: dict,
                          ml_prediction: dict | None) -> tuple:
    agent = Agent(model=model, tools=toolkit.all(), system_prompt=SYSTEM_PROMPT)
    prompt = _build_prompt(student, fired_rules, as_of_date, ids, ml_prediction)
    try:
        response = agent(prompt)
        return _parse_agent_output(str(response), fired_rules)
    except Exception as e:
        logger.exception("Agent call failed for %s: %s", student.get("name"), str(e))
        default_category = fired_rules[0] if fired_rules else FALLBACK_CATEGORY
        return default_category, f"[Agent call raised an exception: {e}]"

# --- Output builder ---
RESEARCH_HEADERS = [
    "Canvas User ID", "Student Name", "NSHE ID", "Course Name", "Canvas Course ID", "Section Number",
    "Fired Rules (context)", "Current Grade (%)", "Quiz Score (%)", "Missing Assignments",
    "Total Clicks", "Click Coverage Z", "Click Intensity Z", "Click Coherence Z", "Last Active",
    "ML At-Risk Predicted", "ML At-Risk Probability",
    "Category (AI+ML)", "Rationale (AI+ML)",
    "Category (AI-only)", "Rationale (AI-only)",
]
COL_WIDTHS = [16, 28, 14, 36, 18, 16, 40, 16, 14, 12, 12, 16, 16, 16, 14, 16, 18, 40, 60, 40, 60]

def build_research_xlsx_bytes(records: list) -> bytes:
    wb = openpyxl.Workbook()

    def _write_sheet(ws, sheet_records):
        ws.append(RESEARCH_HEADERS)
        for cell in ws[1]:
            cell.font = Font(name="Arial", bold=True)
            cell.alignment = Alignment(horizontal="center")

        for rec in sorted(sheet_records, key=lambda r: r["student_name"]):
            ws.append([
                rec["canvas_user_id"], rec["student_name"], rec["nshe_id"], rec["course_name"],
                rec["canvas_course_id"], rec["section_number"],
                " | ".join(rec.get("fired_rules", [])),
                rec.get("current_grade"), rec.get("quiz_score"), rec.get("missing_assignments"),
                rec.get("total_clicks"), rec.get("click_coverage_z"), rec.get("click_intensity_z"),
                rec.get("click_coherence_z"), rec.get("last_active"),
                rec.get("at_risk_predicted"), rec.get("at_risk_probability"),
                rec.get("category_ai_ml"), rec.get("rationale_ai_ml"),
                rec.get("category_ai_only"), rec.get("rationale_ai_only"),
            ])
            ri = ws.max_row
            for col in (8, 9, 12, 13, 14, 17):
                ws.cell(ri, col).number_format = "0.00"
            for col in (19, 21):
                ws.cell(ri, col).alignment = Alignment(wrap_text=True, vertical="top")

        for i, w in enumerate(COL_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb.remove(wb.active)
    by_course = defaultdict(list)
    for rec in records:
        by_course[rec["course_name"]].append(rec)
    for course_name in sorted(by_course):
        _write_sheet(wb.create_sheet(course_name[:31]), by_course[course_name])

    buf = __import__("io").BytesIO()
    wb.save(buf)
    return buf.getvalue()

def _upload(body: bytes, bucket: str, key: str, content_type: str) -> str:
    boto3.client("s3").put_object(Bucket=bucket, Key=key, Body=body, ContentType=content_type)
    url = f"s3://{bucket}/{key}"
    logger.info("Uploaded %d bytes -> %s", len(body), url)
    return url

# --- Main entry point ---
def run_research(mydb_cfg: dict, timescale_cfg: dict, canvas_course_ids: list,
                  s3_bucket: str, as_of_date: str) -> dict:
    """
    Runs every student in canvas_course_ids through both agent conditions
    and writes a single research xlsx to S3. Does NOT touch `alerts`,
    `chiron_runs`, or anything the dashboard reads.
    """
    now = datetime.utcnow()
    ts = now.strftime("%Y%m%d_%H%M%S")
    prefix = f"research/{now.strftime('%Y/%m/%d')}"

    roster = fetch_student_roster(mydb_cfg, canvas_course_ids, as_of_date)
    click_signals = fetch_click_signals(timescale_cfg, canvas_course_ids, as_of_date)
    model_predictions = fetch_model_predictions(mydb_cfg, canvas_course_ids, as_of_date)

    canvas_user_ids = list({s["canvas_user_id"] for s in roster})
    ts_student_map, ts_course_map = resolve_timescale_ids(timescale_cfg, canvas_user_ids, canvas_course_ids)

    toolkit = AttributionToolkit(mydb_cfg, timescale_cfg, as_of_date=as_of_date)
    model = BedrockModel(
        model_id=os.environ.get("BEDROCK_MODEL_ID", "us.anthropic.claude-haiku-4-5-20051001-v1:0"),
        region_name=os.environ.get("AWS_DEFAULT_REGION", "us-west-2"),
        temperature=0.0,
    )

    records = []
    for student in roster:
        key = (student["canvas_user_id"], student["canvas_course_id"])
        click = click_signals.get(key)
        prediction = model_predictions.get(key)
        at_risk_flagged = prediction["at_risk_predicted"] if prediction else False

        fired = get_fired_rules(student, click, at_risk_flagged)

        ids = {
            "mydb_student_id": student["student_id"],
            "mydb_course_id": student["course_id"],
            "ts_student_id": ts_student_map.get(student["canvas_user_id"]),
            "ts_course_id": ts_course_map.get(student["canvas_course_id"]),
        }

        cat_ai_only, rat_ai_only = _run_agent_condition(
            toolkit, model, student, fired, as_of_date, ids, ml_prediction=None
        )
        cat_ai_ml, rat_ai_ml = _run_agent_condition(
            toolkit, model, student, fired, as_of_date, ids, ml_prediction=prediction
        )

        logger.info("Research result: %s -> AI-only=%s | AI+ML=%s", student["name"], cat_ai_only, cat_ai_ml)

        records.append({
            "canvas_user_id": student["canvas_user_id"],
            "student_name": student["name"],
            "nshe_id": student["integration_id"],
            "course_name": student["course_name"],
            "canvas_course_id": student["canvas_course_id"],
            "section_number": student["section_number"],
            "fired_rules": fired,
            "current_grade": float(student["current_score"]) if student["current_score"] else None,
            "quiz_score": float(student["quiz_score"]) if student.get("quiz_score") is not None else None,
            "missing_assignments": int(student["missing_assignments"] or 0),
            "total_clicks": click["total_clicks"] if click else 0,
            "click_coverage_z": click["breadth_z"] if click else None,
            "click_intensity_z": click["intensity_z"] if click else None,
            "click_coherence_z": click["coherence_z"] if click else None,
            "last_active": click["last_active"] if click else None,
            "at_risk_predicted": prediction["at_risk_predicted"] if prediction else None,
            "at_risk_probability": prediction["at_risk_probability"] if prediction else None,
            "category_ai_ml": cat_ai_ml,
            "rationale_ai_ml": rat_ai_ml,
            "category_ai_only": cat_ai_only,
            "rationale_ai_only": rat_ai_only,
        })

    s3_excel = _upload(
        build_research_xlsx_bytes(records), s3_bucket,
        f"{prefix}/research_alerts_{ts}.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    logger.info("Research run complete: %d students, written to %s", len(records), s3_excel)
    return {"total": len(records), "s3_excel": s3_excel, "batch_timestamp": ts, "date_prefix": prefix}

# --- Optional Lambda entry point, invoked separately from the live pipeline ---
def lambda_handler(event, context):
    body = json.loads(event["body"]) if isinstance(event.get("body"), str) else event.get("body", event)
    canvas_course_ids = body["canvas_course_ids"]
    as_of_date = body.get("as_of_date") or datetime.utcnow().strftime("%Y-%m-%d")

    mydb_cfg, timescale_cfg = _build_db_configs()
    bucket = os.environ["ALERTS_BUCKET"]

    result = run_research(mydb_cfg, timescale_cfg, canvas_course_ids, bucket, as_of_date)
    return {"statusCode": 200, "body": json.dumps(result, indent=2)}